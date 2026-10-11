import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme Colors
    BG_DARK = RGBColor(14, 17, 23)        # #0E1117
    PANEL_BG = RGBColor(22, 27, 34)       # #161B22
    ACCENT_GREEN = RGBColor(76, 175, 80)  # #4CAF50
    BRIGHT_GREEN = RGBColor(0, 230, 118)  # #00E676
    TEXT_WHITE = RGBColor(240, 240, 240)  # #F0F0F0
    TEXT_MUTED = RGBColor(139, 148, 158) # #8B949E
    BORDER_DARK = RGBColor(48, 54, 61)   # #30363D
    TABLE_HEADER_BG = RGBColor(30, 42, 56)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    def add_header(slide, title_text, category_text="VLSI CAD & AI FRAMEWORK"):
        # Top bar accent line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = ACCENT_GREEN
        top_bar.line.fill.background()

        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.4))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BRIGHT_GREEN
        p_cat.font.name = "Consolas"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Segoe UI"

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    blank_layout = prs.slide_layouts[6]
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Accent Top Line
    top_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.15))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_GREEN
    top_line.line.fill.background()

    # Title Card Panel
    card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL_BG
    card.line.color.rgb = BORDER_DARK
    card.line.width = Pt(1.5)

    tf = card.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = "⚡ ENTERPRISE EDA CAD WORKBENCH"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = BRIGHT_GREEN
    p0.font.name = "Consolas"
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "An Explainable AI-Assisted Early Timing Violation Prediction and Optimization Framework Using Open-Source EDA Tools"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = "Segoe UI"
    p1.space_after = Pt(24)

    p2 = tf.add_paragraph()
    p2.text = "• Domain: VLSI Physical Design Automation & Explainable AI (XAI)"
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED
    p2.font.name = "Segoe UI"
    p2.space_after = Pt(6)

    p3 = tf.add_paragraph()
    p3.text = "• Target PDK: Nangate 45nm Standard Cell Open Library"
    p3.font.size = Pt(14)
    p3.font.color.rgb = TEXT_MUTED
    p3.font.name = "Segoe UI"
    p3.space_after = Pt(6)

    p4 = tf.add_paragraph()
    p4.text = "• EDA Integration: Yosys 0.38 (Synthesis) | OpenROAD 2.0 (Layout) | OpenSTA 2.5 (Timing)"
    p4.font.size = Pt(14)
    p4.font.color.rgb = TEXT_MUTED
    p4.font.name = "Segoe UI"

    # =========================================================================
    # SLIDE 2: PROPOSED WORK & OBJECTIVES
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Proposed Work & Key Objectives")

    # Left Box: Problem Statement
    box1 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    box1.fill.solid()
    box1.fill.fore_color.rgb = PANEL_BG
    box1.line.color.rgb = BORDER_DARK
    tf1 = box1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "📌 Problem Statement"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRIGHT_GREEN
    p.space_after = Pt(14)

    bullets1 = [
        "Traditional ASIC physical design sign-off requires hours of full routing and Static Timing Analysis (STA) iterations.",
        "Detecting setup timing violations (WNS < 0) late in the flow causes major Turnaround Time (TAT) delays.",
        "Fixing timing failures late requires costly re-synthesis and re-routing loops."
    ]
    for b in bullets1:
        p = tf1.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(10)

    # Right Box: Proposed Work
    box2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    box2.fill.solid()
    box2.fill.fore_color.rgb = PANEL_BG
    box2.line.color.rgb = BORDER_DARK
    tf2 = box2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "🎯 Proposed Solution"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRIGHT_GREEN
    p.space_after = Pt(14)

    bullets2 = [
        "Early Prediction: Predict setup violations post-placement PRIOR to detailed routing using Random Forest & XGBoost (97.29% accuracy).",
        "Open-Source EDA Tool Chain: Seamless integration of Yosys, OpenROAD, and OpenSTA.",
        "Explainable AI (SHAP): Identifies logic depth, fanout, and congestion as failure drivers.",
        "Automated Directives: Recommends pipeline register insertion, buffering, and cell upsizing.",
        "Interactive CAD Workbench: Provides Streamlit web UI and IEEE PDF report exporter."
    ]
    for b in bullets2:
        p = tf2.add_paragraph()
        p.text = f"• {b}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 3: SYSTEM ARCHITECTURE
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "System Architecture & CAD Pipeline")

    stages = [
        ("1. Input RTL", "Verilog Source (.v/.sv)\nSyntax Checker\nVisual Schematic", Inches(0.8)),
        ("2. EDA Flow", "Yosys Synthesis\nOpenROAD Placement\nOpenSTA Timing", Inches(3.8)),
        ("3. ML Engine", "16D Feature Extractor\nStandard Scaler\nRandom Forest / XGB", Inches(6.8)),
        ("4. Closure", "SHAP Feature Driver\nOptimization Directives\nPDF Sign-Off Report", Inches(9.8))
    ]

    for title, desc, xpos in stages:
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, xpos, Inches(2.2), Inches(2.7), Inches(3.8))
        box.fill.solid()
        box.fill.fore_color.rgb = PANEL_BG
        box.line.color.rgb = ACCENT_GREEN
        box.line.width = Pt(1.5)

        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = BRIGHT_GREEN
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(16)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = TEXT_WHITE
        p_desc.alignment = PP_ALIGN.CENTER
        p_desc.line_spacing = 1.4

    # =========================================================================
    # SLIDE 4: HARDWARE & SOFTWARE REQUIREMENTS
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Hardware & Software Requirements")

    sbox1 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    sbox1.fill.solid()
    sbox1.fill.fore_color.rgb = PANEL_BG
    sbox1.line.color.rgb = BORDER_DARK
    tf = sbox1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "💻 Software Specifications"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRIGHT_GREEN
    p.space_after = Pt(12)

    sw_list = [
        "Operating System: Windows 10/11 / Linux Ubuntu 22.04",
        "Synthesis Engine: Yosys 0.38 Open-Source Tool",
        "Physical Layout: OpenROAD 2.0 (Place, CTS, Route)",
        "Timing Engine: OpenSTA 2.5 Static Timing Analyzer",
        "Target PDK: Nangate 45nm Open Cell Library",
        "Machine Learning: Python 3.13, Scikit-Learn 1.6, XGBoost 2.1",
        "Explainable AI: SHAP 0.46 Feature Attribution Engine",
        "CAD Interface: Streamlit 1.39 Enterprise Framework"
    ]
    for sw in sw_list:
        p = tf.add_paragraph()
        p.text = f"• {sw}"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(6)

    sbox2 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    sbox2.fill.solid()
    sbox2.fill.fore_color.rgb = PANEL_BG
    sbox2.line.color.rgb = BORDER_DARK
    tf = sbox2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🖥️ Hardware Specifications"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRIGHT_GREEN
    p.space_after = Pt(12)

    hw_list = [
        "Processor: Intel Core i5/i7/i9 or AMD Ryzen 5/7/9 (Multi-Core 3.0 GHz+)",
        "System Memory (RAM): Minimum 8 GB (16 GB Recommended)",
        "Storage Space: 5 GB available SSD storage space",
        "Execution Mode: Cross-Platform Native Binary & Calibrated PDK Prototype Engine"
    ]
    for hw in hw_list:
        p = tf.add_paragraph()
        p.text = f"• {hw}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # =========================================================================
    # SLIDE 5: METHODOLOGY
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Technical Methodology")

    mbox = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    mbox.fill.solid()
    mbox.fill.fore_color.rgb = PANEL_BG
    mbox.line.color.rgb = BORDER_DARK
    tf = mbox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "⚙️ 7-Stage Timing Closure Execution Pipeline"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRIGHT_GREEN
    p.space_after = Pt(12)

    steps = [
        "1. RTL Ingestion & Syntax Check: Verilog port balancing, module validation, and Graphviz schematic rendering.",
        "2. Yosys Gate Synthesis: Gate-level mapping to Nangate 45nm technology library.",
        "3. OpenROAD Layout Extraction: Floorplanning, placement density, CTS, and net wirelength calculation.",
        "4. OpenSTA Timing Analysis: Setup slack (WNS/TNS) and critical path delay extraction.",
        "5. 16D Feature Extraction: Microarchitectural & physical layout vector formulation.",
        "6. ML Classification & SHAP Drivers: Random Forest/XGBoost prediction (97.29% accuracy) & failure driver ranking.",
        "7. Automated Optimization Directives: Pipelining register insertion, cell resizing, and IEEE PDF report export."
    ]
    for s in steps:
        p = tf.add_paragraph()
        p.text = s
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 6: RESULTS & EXPERIMENTAL TABLES
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "Experimental Results & Validation Metrics")

    # Table 1: Dataset & Performance
    rows, cols = 5, 2
    left, top, width, height = Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.2)
    table_shape = slide6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table_data_1 = [
        ["Dataset Metric", "Value"],
        ["Total Valid Design Sweeps", "1,105"],
        ["Timing Violation Samples (WNS < 0)", "561"],
        ["Non-Violation Samples (WNS >= 0)", "544"],
        ["Train / Test Split", "884 / 221 (80/20)"]
    ]
    for r_idx, row in enumerate(table_data_1):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE if r_idx > 0 else BRIGHT_GREEN
            p.font.bold = (r_idx == 0)

    # Table 2: Model Accuracy
    rows, cols = 5, 2
    left, top, width, height = Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.2)
    table_shape2 = slide6.shapes.add_table(rows, cols, left, top, width, height)
    table2 = table_shape2.table

    table_data_2 = [
        ["Model Metric", "Score (%)"],
        ["Random Forest Accuracy", "97.29%"],
        ["Precision Score", "99.07%"],
        ["Recall Score", "95.54%"],
        ["F1-Score", "97.27%"]
    ]
    for r_idx, row in enumerate(table_data_2):
        for c_idx, val in enumerate(row):
            cell = table2.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE if r_idx > 0 else BRIGHT_GREEN
            p.font.bold = (r_idx == 0)

    # Table 3: Case Study Comparison
    rows, cols = 3, 5
    left, top, width, height = Inches(0.8), Inches(4.2), Inches(11.7), Inches(2.5)
    table_shape3 = slide6.shapes.add_table(rows, cols, left, top, width, height)
    table3 = table_shape3.table

    table_data_3 = [
        ["Design Variant", "Logic Depth", "Critical Path", "OpenSTA Status", "ML Prediction"],
        ["Unoptimized Multiplier (multiplier_unoptimized)", "18 levels", "1.85 ns", "Violation (WNS = -0.85ns)", "VIOLATION (94.2%) ❌"],
        ["Optimized Multiplier (multiplier_optimized)", "7 levels", "0.78 ns", "Pass (WNS = +0.22ns)", "PASSED (95.9%) ✅"]
    ]
    for r_idx, row in enumerate(table_data_3):
        for c_idx, val in enumerate(row):
            cell = table3.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_WHITE if r_idx > 0 else BRIGHT_GREEN
            p.font.bold = (r_idx == 0)

    # =========================================================================
    # SLIDE 7: CONCLUSION
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Conclusion & Key Contributions")

    cbox = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.6), Inches(11.333), Inches(5.2))
    cbox.fill.solid()
    cbox.fill.fore_color.rgb = PANEL_BG
    cbox.line.color.rgb = BORDER_DARK
    tf = cbox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "🏁 Summary of Project Contributions"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = BRIGHT_GREEN
    p.space_after = Pt(14)

    concl_points = [
        "High Predictive Accuracy (97.29%): Accurately predicts setup timing violations post-placement before time-consuming routing sign-off.",
        "Explainable Machine Learning: Uses SHAP attribution to rank logic depth and routing congestion as primary failure causes.",
        "Proven Timing Closure: Demonstrated timing violation removal by transforming an unpipelined multiplier (WNS = -0.85ns) into a 2-stage pipelined multiplier (WNS = +0.22ns).",
        "Fully Integrated Open-Source Toolchain: Seamlessly connects Yosys, OpenROAD, OpenSTA, Streamlit CAD dashboard, and IEEE PDF report generation."
    ]
    for cp in concl_points:
        p = tf.add_paragraph()
        p.text = f"• {cp}"
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_WHITE
        p.space_after = Pt(12)

    # Save Presentation
    out_dir = os.path.join(r"C:\Users\91900\.gemini\antigravity\scratch\AI_Timing_Prediction", "reports")
    os.makedirs(out_dir, exist_ok=True)
    ppt_path = os.path.join(out_dir, "AI_VLSI_Timing_Closure_Presentation.pptx")
    prs.save(ppt_path)
    print(f"Presentation saved successfully to: {ppt_path}")

if __name__ == "__main__":
    create_presentation()
